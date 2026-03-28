import os
import logging
from typing import List, Optional
import PyPDF2
from pptx import Presentation

logger = logging.getLogger(__name__)

class FileProcessor:
    """Provides methods for reading and extracting text from supported file formats."""

    @staticmethod
    def read_txt(file_path: str) -> str:
        """Extracts text from a .txt file."""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            logger.error(f"Error reading TXT {file_path}: {e}")
            return ""

    @staticmethod
    def read_pdf(file_path: str) -> str:
        """Extracts text from a .pdf file."""
        text = ""
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f, strict=False)
                for page in reader.pages:
                    text += page.extract_text() or ""
        except Exception as e:
            logger.error(f"Error reading PDF {file_path}: {e}")
        return text

    @staticmethod
    def read_pptx(file_path: str) -> str:
        """Extracts text from a .pptx file."""
        text = ""
        try:
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        except Exception as e:
            logger.error(f"Error reading PPTX {file_path}: {e}")
        return text

    def extract_text_from_directory(self, directory_path: str) -> str:
        """Reads all supported files in a directory and returns concatenated text."""
        all_text = ""
        for filename in os.listdir(directory_path):
            file_path = os.path.join(directory_path, filename)
            logger.info(f"Processing file: {filename}")

            if filename.lower().endswith(".txt"):
                all_text += self.read_txt(file_path) + "\n\n"
            elif filename.lower().endswith(".pdf"):
                all_text += self.read_pdf(file_path) + "\n\n"
            elif filename.lower().endswith(".pptx"):
                all_text += self.read_pptx(file_path) + "\n\n"

        return all_text
